import pandas as pd
from haversine import haversine
import numpy as np
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
# import matplotlib.pyplot as plt
# import folium
# from pyecharts.charts import Map
# from pyecharts import options as opts
import folium
# from folium.plugins import MarkerCluster

print("正在读取‘报账点信息管理’")
df = pd.read_excel('报账点信息管理.xlsx')
print("读取成功！")
# 读取新基站报价汇总数据
print("正在读取‘新机房报价汇总’")
new_stations_df = pd.read_excel('新机房报价汇总.xlsx')
print("读取成功！")

# # 创建地图对象
# m = Map()

# 数据清理
# 除去机房面积，经度和纬度的缺失值
df_cleaned = df.dropna(subset=['机房面积', '经度', '纬度', '合同年租金'])


# 用户输入给定的比价范围
radius_km = float(input("请输入给定的半径（公里）（例如：1）："))

# 设置存储全部范围内存量基站的全局变量
global_nearby_df = pd.DataFrame()

# 通过经纬度计算距离
def haversine_distance(row, new_lat, new_lon):
    return haversine((row['纬度'], row['经度']), (new_lat, new_lon))


# 在循环之前，先为 df_cleaned 计算所有点到新基站的距离
def preprocess_distances(df_cleaned, new_stations_df):
    # 创建一个新的 DataFrame 来存储距离
    # distances_df = pd.DataFrame(index=df_cleaned.index)
    distances_df = df_cleaned.copy()

    for index, new_station in new_stations_df.iterrows():
        new_lat = new_station['纬度']
        new_lon = new_station['经度']
        distances_df[f'距离_{index+2}'] = df_cleaned.apply(
            lambda row: haversine_distance(row, new_lat, new_lon),
            axis=1
        )
    return distances_df


# 计算距离
distances_df = preprocess_distances(df_cleaned, new_stations_df)
# distances_df.head(5)

# test
# distances_df.to_excel('ceshi.xlsx', index=False)

# 定义一个函数来处理每个新基站的数据（不需要重新计算距离）
def process_new_station(row, distances_df, radius_km):
    new_lat = row['纬度']
    new_lon = row['经度']
    new_rental = row['合同年租金']
    global global_nearby_df
    # 使用预先计算的距离

    nearby_distances = distances_df[f'距离_{row.name+2}']
    # nearby_distances.to_excel('车市.xlsx', index=False)
    # 筛选范围内存量机房的索引值
    filtered_indices = [i for i, value in enumerate(nearby_distances) if value <= radius_km]
    nearby_df = df_cleaned.iloc[filtered_indices]
    # 删除大于阈值的行
    nearby_df = nearby_df[nearby_df['合同年租金'] < 30000]
    # 筛选nearby_distances
    filtered_elements = nearby_distances.iloc[filtered_indices]
    # nearby_df.loc[:, f'距离_{row.name + 2}'] = nearby_distances.iloc[filtered_indices]
    # 给nearby_df后面加一列距离
    nearby_df['距离'] = np.nan
    nearby_df['距离'] = filtered_elements

    nearby_df.to_excel(f'新机房_{row.name+2}周边的存量机房汇总.xlsx', index=False)
    # 计算平均价格、最大值、最小值
    avg_rental = nearby_df['合同年租金'].mean()
    min_rental = nearby_df['合同年租金'].min()
    max_rental = nearby_df['合同年租金'].max()
    # nearest_rental = nearby_df['合同年租金'].iloc[0] if not nearby_df.empty else None
    #
    # 筛选出非零的'距离'值
    non_zero_distances = nearby_df[nearby_df['距离'] != 0]

    # 如果非零的'距离'存在，找到最小的那个，并获取对应的'合同年租金'
    if not non_zero_distances.empty:
        min_distance_idx = non_zero_distances['距离'].idxmin()  # 获取最小'距离'值的索引
        nearest_rental = non_zero_distances.at[min_distance_idx, '合同年租金']  # 使用索引获取'合同年租金'
    else:
        nearest_rental = None  # 如果没有非零的'距离'，则设置为None
    # nearest_rental = nearby_df.at[nearby_df['距离'].idxmin(), '合同年租金']

    # 给全局变量赋值
    # 如果global_df为空，则直接赋值
    if global_nearby_df.empty:
        global_nearby_df = nearby_df.copy()
    else:
        # 否则，使用pd.concat追加
        global_nearby_df = pd.concat([global_nearby_df, nearby_df], ignore_index=True)
        # 返回结果
    return {
        '新基站经度': new_lon,
        '新基站纬度': new_lat,
        '新基站合同年租金': new_rental,
        '周围存量机房平均合同年租金': avg_rental,
        '周围存量机房最小合同年租金': min_rental,
        '周围存量机房最大合同年租金': max_rental,
        '最近存量机房合同年租金': nearest_rental
    }

# 复制new_stations_quotes到results_df
results_df = new_stations_df.copy()

# 添加新列，初始值为NaN
new_columns = ['周围存量机房平均合同年租金', '周围存量机房最小合同年租金', '周围存量机房最大合同年租金', '最近存量机房合同年租金']
for column in new_columns:
    results_df[column] = np.nan

for index, row in results_df.iterrows():
    # 这里假设process_new_station可以直接处理row
    # 注意：在实际应用中，可能需要将row转换为process_new_station所需的格式
    rental_data = process_new_station(row, distances_df, radius_km)  # 这里可能需要调整以适应process_new_station的实际输入

    # 更新DataFrame中的列
    results_df.at[index, '周围存量机房平均合同年租金'] = rental_data['周围存量机房平均合同年租金']
    results_df.at[index, '周围存量机房最小合同年租金'] = rental_data['周围存量机房最小合同年租金']
    results_df.at[index, '周围存量机房最大合同年租金'] = rental_data['周围存量机房最大合同年租金']
    results_df.at[index, '最近存量机房合同年租金'] = rental_data['最近存量机房合同年租金']
results_df['新机房租金与周边平均租金对比结果'] = ['<=' if a <= b else '>' for a, b in zip(results_df['合同年租金'], results_df['周围存量机房平均合同年租金'])]
results_df['新机房租金与最近存量机房租金对比结果'] = ['<=' if a <= b else '>' for a, b in zip(results_df['合同年租金'], results_df['最近存量机房合同年租金'])]

results_df.to_excel('新机房周边的机房数据对比.xlsx', index=False)
# global_nearby_df.to_excel('teste.xlsx', index=False)
# 给新机房周边的机房数据对比.xlsx打标识
wb = load_workbook('新机房周边的机房数据对比.xlsx')
ws = wb.active

# 遍历除了标题行之外的所有行
for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=False), start=2):
    # 获取'对比结果'列的单元格
    cell_result_1 = row[results_df.columns.get_loc('新机房租金与周边平均租金对比结果')]
    cell_result_2 = row[results_df.columns.get_loc('新机房租金与最近存量机房租金对比结果')]

    # 根据'对比结果'的值来设置单元格的背景色
    if cell_result_1.value == '<=':
        cell_result_1.fill = PatternFill(start_color="90EE90", end_color="90EE90", fill_type="solid")
    else:
        cell_result_1.fill = PatternFill(start_color="FF6347", end_color="FF6347", fill_type="solid")

    if cell_result_2.value == '<=':
        cell_result_2.fill = PatternFill(start_color="90EE90", end_color="90EE90", fill_type="solid")
    else:
        cell_result_2.fill = PatternFill(start_color="FF6347", end_color="FF6347", fill_type="solid")

    # 保存修改后的Excel文件
wb.save('新机房周边的机房数据对比.xlsx')



print("处理完成，结果已保存到 '新机房周边的机房数据对比.xlsx'")


m = folium.Map(location=[34.263161, 108.948022], zoom_start=7)

# 添加Global Nearby点的蓝色标记
for index, row in global_nearby_df.iterrows():
    # readable_index = int(index) + 2
    folium.Marker(
        location=[row['纬度'], row['经度']],  # 注意folium的经纬度顺序是[纬度, 经度]
        popup=f'合同年租金: {row["合同年租金"]}元',
        # popup=f'行号: {readable_index}, 合同年租金: {row["合同年租金"]}元',
        icon=folium.Icon(color='blue')
    ).add_to(m)

# 添加New Stations点的红色标记
for index, row in new_stations_df.iterrows():
    readable_index = int(index) + 2
    folium.Marker(
        location=[row['纬度'], row['经度']],
        # popup=f'合同年租金: {row["合同年租金"]}元',
        popup=f'行号: {readable_index}, 合同年租金: {row["合同年租金"]}元',
        icon=folium.Icon(color='red')
    ).add_to(m)

# 保存地图为HTML文件
m.save('新增与存量机房地图信息.html')

# 如果你在Jupyter Notebook中工作，可以使用下面的代码来直接在Notebook中显示地图
# from IPython.display import IFrame
# IFrame(src='map.html', width=700, height=450)

